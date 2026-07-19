"""document_image_service — 从 Office 文档中提取嵌入图片

支持:
  - .docx (python-docx)
  - .xlsx (openpyxl)

输出统一为 multimodal 可用的 base64 图片字典，复用 file_service 的压缩逻辑。
"""

import base64
import io
import logging
from pathlib import Path
from typing import Any

from app.config import MAX_IMAGES_PER_REQUEST
from app.services.file_service import encode_image_to_base64

logger = logging.getLogger("gt_agent.document_image")

# Office 文档单次提取上限（避免一次性 token 爆炸）
MAX_DOC_IMAGES = MAX_IMAGES_PER_REQUEST * 4  # 通常比直接上传图片稍宽松

_IMAGE_CONTENT_TYPES = {
    "image/png",
    "image/jpeg",
    "image/jpg",
    "image/gif",
    "image/bmp",
    "image/webp",
}


def _is_image_part(content_type: str | None) -> bool:
    if not content_type:
        return False
    return content_type.lower() in _IMAGE_CONTENT_TYPES


def _to_data_uri(image_info: dict) -> str:
    return f"data:{image_info['mime']};base64,{image_info['data']}"


def _extract_docx_images(file_content: bytes, doc_name: str) -> list[dict]:
    """从 docx 中抽取 inline_shapes 与 package parts 里的图片。"""
    try:
        from docx import Document
    except ImportError:
        logger.warning("python-docx 未安装，无法提取 Word 图片")
        return []

    images: list[dict] = []
    try:
        doc = Document(io.BytesIO(file_content))
    except Exception as e:
        logger.warning("document_image: 无法打开 docx %s: %s", doc_name, e)
        return images

    # 1) inline_shapes（最常见：段落中嵌入的图片）
    # python-docx 的 InlineShape 没有 .part 属性，需通过 blip 的 rId 取 related_parts
    for idx, shape in enumerate(doc.inline_shapes):
        try:
            blip = shape._inline.graphic.graphicData.pic.blipFill.blip
            rId = blip.embed
            if not rId:
                continue
            image_part = doc.part.related_parts.get(rId)
            if not image_part or not _is_image_part(getattr(image_part, "content_type", None)):
                continue
            img_bytes = image_part.blob
            if not img_bytes:
                continue
            ext = _mime_to_ext(getattr(image_part, "content_type", None))
            img_filename = f"{doc_name}_inline_{idx + 1}{ext}"
            info = encode_image_to_base64(img_bytes, img_filename)
            info["source"] = "docx_inline_shape"
            images.append(info)
        except Exception as e:
            logger.debug("document_image: 提取 docx inline_shape %d 失败: %s", idx, e)

    # 2) 遍历所有 package parts，兜底提取未关联到 inline_shapes 的图片
    existing_hashes = {hash(img["data"]) for img in images}  # base64 data 的 hash
    try:
        for idx, part in enumerate(doc.part.package.parts):
            try:
                if not _is_image_part(part.content_type):
                    continue
                img_bytes = part.blob
                if not img_bytes:
                    continue
                info = encode_image_to_base64(img_bytes, f"{doc_name}_part_{idx + 1}")
                # 基于压缩/编码后的 base64 data 去重
                if hash(info["data"]) in existing_hashes:
                    continue
                existing_hashes.add(hash(info["data"]))
                info["filename"] = f"{doc_name}_part_{idx + 1}{_mime_to_ext(part.content_type)}"
                info["source"] = "docx_part"
                images.append(info)
            except Exception as e:
                logger.debug("document_image: 提取 docx part 失败: %s", e)
    except Exception as e:
        logger.debug("document_image: 遍历 docx parts 失败: %s", e)

    return images


def _extract_xlsx_images(file_content: bytes, sheet_name: str) -> list[dict]:
    """从 xlsx 中抽取所有工作表里的图片。"""
    try:
        from openpyxl import load_workbook
    except ImportError:
        logger.warning("openpyxl 未安装，无法提取 Excel 图片")
        return []

    images: list[dict] = []
    wb = None
    try:
        wb = load_workbook(io.BytesIO(file_content), data_only=True)
    except Exception as e:
        logger.warning("document_image: 无法打开 xlsx %s: %s", sheet_name, e)
        return images

    try:
        for s_idx, ws in enumerate(wb.worksheets, start=1):
            try:
                # openpyxl 2.6+ 使用 _images 存储图片对象
                ws_images = getattr(ws, "_images", [])
                for i_idx, img in enumerate(ws_images, start=1):
                    try:
                        # Image 对象的 _data() 返回 bytes
                        img_bytes = img._data()
                        if not img_bytes:
                            continue
                        ext = _mime_to_ext(img.format) if getattr(img, "format", None) else ".png"
                        img_filename = f"{sheet_name}_sheet{s_idx}_img{i_idx}{ext}"
                        info = encode_image_to_base64(img_bytes, img_filename)
                        info["source"] = "xlsx_image"
                        images.append(info)
                    except Exception as e:
                        logger.debug("document_image: 提取 xlsx 图片失败: %s", e)
            except Exception as e:
                logger.debug("document_image: 遍历 xlsx 工作表失败: %s", e)
    finally:
        try:
            wb.close()
        except Exception:
            pass

    return images


def _mime_to_ext(mime: str | None) -> str:
    if not mime:
        return ".png"
    mapping = {
        "image/png": ".png",
        "image/jpeg": ".jpg",
        "image/jpg": ".jpg",
        "image/gif": ".gif",
        "image/bmp": ".bmp",
        "image/webp": ".webp",
    }
    return mapping.get(mime.lower(), ".png")


def is_office_document(filename: str) -> bool:
    """判断文件是否为支持的 Office 文档（含嵌入图片可被提取）。"""
    return filename.lower().endswith((".docx", ".xlsx"))


def extract_document_images(file_content: bytes, filename: str, max_images: int = MAX_DOC_IMAGES) -> dict:
    """从 Office 文档中提取嵌入图片。

    Args:
        file_content: 文件二进制内容
        filename: 原始文件名（用于生成图片别名）
        max_images: 最多返回多少张图片，超过则截断

    Returns:
        {
            "filename": str,
            "total": int,          # 实际提取到的数量
            "returned": int,       # 返回的数量（受 max_images 限制）
            "truncated": bool,
            "images": [            # 每个元素与 encode_image_to_base64 返回一致
                {"mime": "image/jpeg", "data": "base64...", "filename": "...", ...},
            ],
        }
    """
    lower_name = filename.lower()
    base_name = Path(filename).stem

    images: list[dict] = []
    if lower_name.endswith(".docx"):
        images = _extract_docx_images(file_content, base_name)
    elif lower_name.endswith(".xlsx"):
        images = _extract_xlsx_images(file_content, base_name)
    else:
        return {
            "filename": filename,
            "total": 0,
            "returned": 0,
            "truncated": False,
            "images": [],
            "unsupported": True,
        }

    total = len(images)
    truncated = total > max_images
    if truncated:
        images = images[:max_images]
        logger.info(
            "document_image: %s 提取到 %d 张图片，已截断至 %d 张",
            filename, total, max_images,
        )
    else:
        logger.info("document_image: %s 提取到 %d 张图片", filename, total)

    return {
        "filename": filename,
        "total": total,
        "returned": len(images),
        "truncated": truncated,
        "images": images,
    }


def images_to_content_blocks(images: list[dict]) -> list[dict[str, Any]]:
    """将 extract_document_images 返回的图片列表转为 OpenAI multimodal content blocks。"""
    blocks = []
    for img in images:
        blocks.append({
            "type": "image_url",
            "image_url": {"url": _to_data_uri(img)},
        })
    return blocks
